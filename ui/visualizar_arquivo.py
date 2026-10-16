"""Pré-visualização de arquivos (gerados pela ÁGORA ou anexados) numa janela
da própria tela, antes de baixar.

O arquivo é baixado do Storage pelo cliente da pessoa logada, como no botão
Baixar, e montado aqui mesmo: nada é enviado a serviços de fora (visualizador
do Office, Google etc.). PDF vira imagem página a página; slides, Word e
planilhas são lidos e mostrados com o visual da Latitudes.
"""

from collections.abc import Callable
import csv
from html import escape
from io import BytesIO, StringIO

import streamlit as st

PREVIEW_CACHE_KEY = "file_preview_cache"
PREVIEW_CACHE_MAX_ITEMS = 3
MAX_PDF_PAGES = 15
MAX_SLIDES = 40
MAX_TABLE_ROWS = 500
PDF_RENDER_SCALE = 1.4

PPTX_MIME_TYPE = (
    "application/vnd.openxmlformats-officedocument.presentationml.presentation"
)
DOCX_MIME_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)
XLSX_MIME_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
PREVIEW_MIME_TYPES = {
    "application/pdf",
    PPTX_MIME_TYPE,
    DOCX_MIME_TYPE,
    XLSX_MIME_TYPE,
    "text/csv",
    "text/plain",
}


def can_preview(mime_type: str) -> bool:
    return mime_type in PREVIEW_MIME_TYPES


def _load_cached(
    path: str,
    load_file: Callable[[], bytes],
) -> bytes:
    """Guarda os últimos arquivos abertos na sessão: a janela roda de novo a
    cada interação e o arquivo não muda."""
    cache = st.session_state.setdefault(PREVIEW_CACHE_KEY, {})

    if path not in cache:
        if len(cache) >= PREVIEW_CACHE_MAX_ITEMS:
            cache.pop(next(iter(cache)))

        cache[path] = load_file()

    return cache[path]


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------


def _show_pdf(data: bytes) -> None:
    import pypdfium2

    document = pypdfium2.PdfDocument(data)

    try:
        total = len(document)

        for index in range(min(total, MAX_PDF_PAGES)):
            page = document[index]
            image = page.render(scale=PDF_RENDER_SCALE).to_pil()
            st.image(
                image,
                width="stretch",
            )
            page.close()
    finally:
        document.close()

    if total > MAX_PDF_PAGES:
        st.caption(
            f"Mostrando as {MAX_PDF_PAGES} primeiras de {total} páginas. "
            "Baixe o arquivo para ver tudo."
        )


# --------------------------------------------------------------------------
# Slides
# --------------------------------------------------------------------------


def _slide_background(slide) -> str:
    try:
        fill = slide.background.fill
        return str(fill.fore_color.rgb)
    except Exception:
        return "FFFFFF"


def _is_dark(color: str) -> bool:
    try:
        red, green, blue = (int(color[index : index + 2], 16) for index in (0, 2, 4))
    except ValueError:
        return False

    return (red * 299 + green * 587 + blue * 114) / 1000 < 160


def _line_html(text: str) -> str:
    """Linha do slide; o marcador ("•\\t" ou "1.\\t") ganha destaque."""
    marker, separator, rest = text.partition("\t")

    if separator and len(marker) <= 4:
        return (
            f'<div class="pv-line"><span class="pv-marker">{escape(marker)}</span>'
            f"{escape(rest)}</div>"
        )

    return f'<div class="pv-line">{escape(text)}</div>'


def _table_html(rows: list[list[str]]) -> str:
    if not rows:
        return ""

    header, *body = rows
    head = "".join(f"<th>{escape(cell)}</th>" for cell in header)
    lines = "".join(
        "<tr>" + "".join(f"<td>{escape(cell)}</td>" for cell in row) + "</tr>"
        for row in body
    )
    return f'<table class="pv-table"><tr>{head}</tr>{lines}</table>'


def _show_slides(data: bytes) -> None:
    from pptx import Presentation

    presentation = Presentation(BytesIO(data))
    slides = list(presentation.slides)
    parts = []

    for number, slide in enumerate(slides[:MAX_SLIDES], start=1):
        frames: list[list[str]] = []
        tables: list[list[list[str]]] = []

        for shape in slide.shapes:
            if getattr(shape, "has_table", False) and shape.has_table:
                tables.append(
                    [
                        [cell.text for cell in row.cells]
                        for row in shape.table.rows
                    ]
                )
            elif shape.has_text_frame and shape.text_frame.text.strip():
                lines = [
                    paragraph.text
                    for paragraph in shape.text_frame.paragraphs
                    if paragraph.text.strip()
                ]

                # Número da página no rodapé: não é conteúdo.
                if len(lines) == 1 and lines[0].strip().isdigit():
                    continue

                frames.append(lines)

        background = _slide_background(slide)
        dark = _is_dark(background)

        if number == 1:
            kind = "pv-cover"
        elif dark and not tables and len(frames) == 1:
            kind = "pv-section"
        else:
            kind = ""

        theme = "pv-dark" if dark else ""
        title = frames[0][0] if frames else ""
        body_lines = (frames[0][1:] if frames else []) + [
            line for frame in frames[1:] for line in frame
        ]
        body = "".join(_line_html(line) for line in body_lines)
        body += "".join(_table_html(rows) for rows in tables)

        parts.append(
            f'<div class="pv-slide {kind} {theme}" style="background:#{escape(background)}">'
            f'<div class="pv-slide-title">{escape(title)}</div>'
            f'<div class="pv-slide-body">{body}</div>'
            f'<div class="pv-slide-number">{number}</div>'
            "</div>"
        )

    st.html("".join(parts))

    if len(slides) > MAX_SLIDES:
        st.caption(
            f"Mostrando os {MAX_SLIDES} primeiros de {len(slides)} slides."
        )

    st.caption(
        "Prévia simplificada. No PowerPoint ou no Google Slides o arquivo "
        "aparece com o logo e o desenho completo."
    )


# --------------------------------------------------------------------------
# Word
# --------------------------------------------------------------------------


def _runs_html(paragraph) -> str:
    parts = []

    for run in paragraph.runs:
        text = escape(run.text)

        if run.bold:
            text = f"<strong>{text}</strong>"

        if run.italic:
            text = f"<em>{text}</em>"

        parts.append(text)

    return "".join(parts)


def _show_docx(data: bytes) -> None:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(BytesIO(data))
    parts = []

    for element in document.element.body.iterchildren():
        tag = element.tag.rsplit("}", 1)[-1]

        if tag == "p":
            paragraph = Paragraph(element, document)

            if not paragraph.text.strip():
                continue

            style = (paragraph.style.name if paragraph.style is not None else "").lower()
            sizes = [run.font.size.pt for run in paragraph.runs if run.font.size]

            # Título do documento da ÁGORA: parágrafo com letra de 20 pt ou mais.
            is_title = style == "title" or bool(sizes and max(sizes) >= 20)

            if is_title:
                parts.append(f"<h1>{escape(paragraph.text)}</h1>")
            elif style.startswith("heading"):
                digit = style[-1:]
                level = min(int(digit) + 1, 4) if digit.isdigit() else 2
                parts.append(f"<h{level}>{escape(paragraph.text)}</h{level}>")
            else:
                parts.append(f"<p>{_runs_html(paragraph)}</p>")
        elif tag == "tbl":
            table = Table(element, document)
            parts.append(
                _table_html(
                    [[cell.text for cell in row.cells] for row in table.rows]
                )
            )

    st.html(f'<div class="pv-page">{"".join(parts)}</div>')


# --------------------------------------------------------------------------
# Planilhas e texto
# --------------------------------------------------------------------------


def _rows_to_table(rows: list[list]) -> None:
    import pandas

    rows = [
        ["" if value is None else value for value in row]
        for row in rows
        if any(value not in (None, "") for value in row)
    ]

    # Cabeçalho = primeira linha com pelo menos duas colunas preenchidas
    # (as planilhas da ÁGORA têm logo e título acima da tabela).
    header_index = next(
        (
            index
            for index, row in enumerate(rows)
            if sum(1 for value in row if value != "") >= 2
        ),
        None,
    )

    if header_index is None:
        st.caption("Sem dados para mostrar.")
        return

    for row in rows[:header_index]:
        text = " ".join(str(value) for value in row if value != "")

        if text:
            st.markdown(f"**{escape(text)}**")

    header = [str(value) or f"Coluna {index + 1}" for index, value in enumerate(rows[header_index])]
    width = len(header)
    body = [
        (list(row) + [""] * width)[:width]
        for row in rows[header_index + 1 : header_index + 1 + MAX_TABLE_ROWS]
    ]
    # Nomes de coluna repetidos quebram o DataFrame.
    seen: dict[str, int] = {}

    for index, name in enumerate(header):
        seen[name] = seen.get(name, 0) + 1

        if seen[name] > 1:
            header[index] = f"{name} ({seen[name]})"

    st.dataframe(
        pandas.DataFrame(body, columns=header).astype(str),
        hide_index=True,
        width="stretch",
    )


def _show_xlsx(data: bytes) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(
        BytesIO(data),
        read_only=True,
        data_only=True,
    )

    try:
        sheets = workbook.worksheets
        tabs = st.tabs([sheet.title for sheet in sheets]) if len(sheets) > 1 else [st.container()]

        for tab, sheet in zip(tabs, sheets):
            with tab:
                _rows_to_table([list(row) for row in sheet.iter_rows(values_only=True)])
    finally:
        workbook.close()


def _show_csv(data: bytes) -> None:
    text = data.decode("utf-8-sig", errors="replace")
    delimiter = ";" if text.count(";") >= text.count(",") else ","
    _rows_to_table(list(csv.reader(StringIO(text), delimiter=delimiter)))


def _show_text(data: bytes) -> None:
    st.text(data.decode("utf-8-sig", errors="replace")[:20000])


VIEWERS = {
    "application/pdf": _show_pdf,
    PPTX_MIME_TYPE: _show_slides,
    DOCX_MIME_TYPE: _show_docx,
    XLSX_MIME_TYPE: _show_xlsx,
    "text/csv": _show_csv,
    "text/plain": _show_text,
}


@st.dialog("Visualizar arquivo", width="large")
def show_preview(
    name: str,
    mime_type: str,
    path: str,
    load_file: Callable[[], bytes],
) -> None:
    st.html(f'<div class="pv-file-name">{escape(name)}</div>')

    try:
        data = _load_cached(
            path=path,
            load_file=load_file,
        )
    except Exception as error:
        print(f"[ARQUIVO] Falha ao abrir para visualizar: {type(error).__name__}", flush=True)
        st.error("Não foi possível abrir o arquivo agora. Tente baixar.")
        return

    try:
        VIEWERS[mime_type](data)
    except Exception as error:
        print(f"[ARQUIVO] Falha na pré-visualização ({mime_type}): {type(error).__name__}", flush=True)
        st.warning("Não foi possível mostrar a prévia deste arquivo. Baixe para abrir.")

    st.download_button(
        "Baixar arquivo",
        data=data,
        file_name=name,
        mime=mime_type,
        icon=":material/download:",
        key=f"preview_download_{path}",
        on_click="ignore",
        type="primary",
    )
