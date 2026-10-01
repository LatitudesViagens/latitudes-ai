"""Exporta respostas da ÁGORA (Markdown) para PDF, Word, Excel e CSV.

Os documentos levam a identidade visual da Latitudes. A conversão é feita
localmente, sem chamar modelos de IA.
"""

import csv
from dataclasses import dataclass, field
from datetime import datetime, timedelta, timezone
from io import BytesIO, StringIO
from pathlib import Path
import re
import unicodedata
from xml.sax.saxutils import escape as xml_escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from markdown_it import MarkdownIt
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlsxImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image as PilImage
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.lib.utils import ImageReader
from reportlab.platypus import (
    HRFlowable,
    Paragraph,
    Preformatted,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

PROJECT_ROOT = Path(__file__).resolve().parents[1]
LOGO_FILE = PROJECT_ROOT / "assets" / "logo-latitudes.png"

# Mesmo fuso de services/agent_runner.py (não importado para evitar ciclo:
# o agente usa este módulo para gerar arquivos).
BRAZIL_TIMEZONE = timezone(
    timedelta(hours=-3),
    name="America/Sao_Paulo",
)

# Identidade visual (mesmas cores de assets/styles.css).
OLIVE = "94825D"
OLIVE_DARK = "625234"
GOLD = "FEBA12"
IVORY = "F6F4F0"
BORDER = "DED8CA"
TEXT = "333333"
MUTED = "9B9B9C"

SIGNATURE = "Latitudes"

_MARKDOWN = MarkdownIt(
    "commonmark",
    {"html": False},
).enable(["table", "strikethrough"])


# --------------------------------------------------------------------------
# Leitura do Markdown
# --------------------------------------------------------------------------


@dataclass
class Run:
    text: str
    bold: bool = False
    italic: bool = False
    link: str | None = None


@dataclass
class Block:
    kind: str
    runs: list[Run] = field(default_factory=list)
    level: int = 0
    depth: int = 0
    marker: str = ""
    rows: list[list[list[Run]]] = field(default_factory=list)
    text: str = ""


def _parse_inline(token) -> list[Run]:
    runs: list[Run] = []
    bold = italic = False
    link = None

    for child in token.children or []:
        if child.type == "strong_open":
            bold = True
        elif child.type == "strong_close":
            bold = False
        elif child.type == "em_open":
            italic = True
        elif child.type == "em_close":
            italic = False
        elif child.type == "link_open":
            link = str(child.attrGet("href") or "")
        elif child.type == "link_close":
            link = None
        elif child.type in {"text", "code_inline"}:
            runs.append(Run(child.content, bold, italic, link))
        elif child.type == "softbreak":
            runs.append(Run(" ", bold, italic, link))
        elif child.type == "hardbreak":
            runs.append(Run("\n", bold, italic, link))
        elif child.type == "image":
            runs.append(Run(child.content or "", bold, italic, link))

    return runs


def parse_markdown(content: str) -> list[Block]:
    tokens = _MARKDOWN.parse(content or "")
    blocks: list[Block] = []
    lists: list[dict] = []
    item_has_text = False
    in_quote = False
    table_rows: list[list[list[Run]]] | None = None
    heading_level = 0

    for token in tokens:
        kind = token.type

        if kind in {"bullet_list_open", "ordered_list_open"}:
            start = int(token.attrGet("start") or 1)
            lists.append(
                {
                    "ordered": kind == "ordered_list_open",
                    "number": start - 1,
                }
            )
        elif kind in {"bullet_list_close", "ordered_list_close"}:
            lists.pop()
        elif kind == "list_item_open":
            lists[-1]["number"] += 1
            item_has_text = False
        elif kind == "heading_open":
            heading_level = int(token.tag[1])
        elif kind == "heading_close":
            heading_level = 0
        elif kind == "blockquote_open":
            in_quote = True
        elif kind == "blockquote_close":
            in_quote = False
        elif kind == "table_open":
            table_rows = []
        elif kind == "table_close":
            if table_rows:
                blocks.append(Block("table", rows=table_rows))
            table_rows = None
        elif kind == "tr_open" and table_rows is not None:
            table_rows.append([])
        elif kind == "hr":
            blocks.append(Block("rule"))
        elif kind in {"fence", "code_block"}:
            blocks.append(Block("code", text=token.content.rstrip()))
        elif kind == "inline":
            runs = _parse_inline(token)

            if table_rows is not None:
                table_rows[-1].append(runs)
            elif heading_level:
                blocks.append(Block("heading", runs=runs, level=heading_level))
            elif lists:
                current = lists[-1]
                marker = ""

                if not item_has_text:
                    marker = (
                        f"{current['number']}."
                        if current["ordered"]
                        else "•"
                    )
                    item_has_text = True

                blocks.append(
                    Block(
                        "list_item",
                        runs=runs,
                        depth=len(lists) - 1,
                        marker=marker,
                    )
                )
            else:
                if in_quote:
                    runs = [
                        Run(run.text, run.bold, True, run.link)
                        for run in runs
                    ]

                blocks.append(Block("paragraph", runs=runs))

    return blocks


# Documentos levam só o roteiro: recomendações, observações, dicas, notas e
# fontes ficam na conversa. Títulos que começam com estas palavras abrem uma
# seção que é removida do documento.
_ADVICE_TITLE_PATTERN = re.compile(
    r"^(?:recomenda|observa|obs\b|dica|sugest|nota|importante|aten[çc][ãa]o|"
    r"fontes?\b|pr[óo]ximos passos|considera[çc]|cuidados)",
    flags=re.IGNORECASE,
)
_HEADING_LINE = re.compile(r"^(#{1,6})\s+(.*)$")
# Linha inteira em negrito funcionando como título: "**Observações:**".
_BOLD_TITLE_LINE = re.compile(r"^\*\*([^*]+?)\*\*:?\s*$")
_NOTE_PARAGRAPH = re.compile(
    r"^[*_]*(?:nota|obs|observa[çc][ãa]o|importante|aten[çc][ãa]o)\b",
    flags=re.IGNORECASE,
)


def _clean_title_text(text: str) -> str:
    return re.sub(r"[*_`:]", "", text).strip()


def remove_advice_sections(content: str) -> str:
    """Remove seções de recomendações/observações/fontes de um Markdown."""
    kept_lines = []
    skipping_level = None

    for line in (content or "").splitlines():
        stripped = line.strip()
        heading = _HEADING_LINE.match(stripped)
        bold_title = _BOLD_TITLE_LINE.match(stripped)

        if heading:
            level = len(heading.group(1))
            title = _clean_title_text(heading.group(2))
        elif bold_title:
            level = 7
            title = _clean_title_text(bold_title.group(1))
        else:
            level = None
            title = ""

        if skipping_level is not None:
            ends_section = (
                level is not None and level <= skipping_level
            ) or stripped in {"---", "***", "___"}

            if not ends_section:
                continue

            skipping_level = None

        if level is not None and _ADVICE_TITLE_PATTERN.match(title):
            skipping_level = level
            continue

        kept_lines.append(line)

    paragraphs = re.split(r"\n\s*\n", "\n".join(kept_lines))
    paragraphs = [
        paragraph
        for paragraph in paragraphs
        if not _NOTE_PARAGRAPH.match(paragraph.strip())
    ]

    cleaned = "\n\n".join(paragraph.strip("\n") for paragraph in paragraphs)
    # Remove linhas horizontais soltas que sobram no fim do documento.
    cleaned = re.sub(r"(?:\n\s*(?:---|\*\*\*|___)\s*)+$", "", cleaned.strip())
    return cleaned.strip() or content


def _plain_text(runs: list[Run]) -> str:
    return "".join(run.text for run in runs).strip()


def _split_title(blocks: list[Block], title: str) -> tuple[str, list[Block]]:
    # Se o conteúdo já começa com um título, ele vira o título do documento
    # (evita "Roteiro de Lisboa" duas vezes seguidas).
    if blocks and blocks[0].kind == "heading" and blocks[0].level <= 2:
        heading = _plain_text(blocks[0].runs)

        if heading:
            return heading, blocks[1:]

    return title, blocks


_DAY_PATTERN = re.compile(r"\bdias?\s+(\d{1,2})\b", flags=re.IGNORECASE)


def looks_like_itinerary(content: str) -> bool:
    """Roteiro = conteúdo organizado por dias (Dia 1, Dia 2...)."""
    days = {int(day) for day in _DAY_PATTERN.findall(content or "")}
    return len(days) >= 2


def looks_like_document(content: str) -> bool:
    """Conteúdo estruturado o bastante para virar PDF/Word."""
    if looks_like_itinerary(content):
        return True

    blocks = parse_markdown(content)

    if any(block.kind == "table" for block in blocks):
        return True

    headings = sum(block.kind == "heading" for block in blocks)
    list_items = sum(block.kind == "list_item" for block in blocks)
    words = len((content or "").split())

    return words >= 120 and (headings >= 2 or list_items >= 4)


def has_table(content: str) -> bool:
    return any(block.kind == "table" for block in parse_markdown(content))


def export_file_name(title: str, extension: str) -> str:
    normalized = unicodedata.normalize("NFKD", title or "")
    ascii_title = normalized.encode("ascii", "ignore").decode("ascii")
    slug = re.sub(r"[^A-Za-z0-9]+", "-", ascii_title).strip("-").lower()
    return f"{slug or 'agora-latitudes'}.{extension}"


def _today() -> str:
    return datetime.now(BRAZIL_TIMEZONE).strftime("%d/%m/%Y")


# --------------------------------------------------------------------------
# PDF (reportlab)
# --------------------------------------------------------------------------

# As fontes padrão do PDF (Helvetica/Times) só cobrem o alfabeto latino
# (cp1252). Símbolos fora dele viram equivalentes simples.
_PDF_REPLACEMENTS = {
    "→": "->",
    "←": "<-",
    "⇒": "=>",
    "✓": "-",
    "✔": "-",
    "✅": "-",
    "❌": "x",
    "★": "*",
    "≈": "~",
    "≥": ">=",
    "≤": "<=",
    " ": " ",
}


def _pdf_safe(text: str) -> str:
    safe_characters = []

    for character in text:
        character = _PDF_REPLACEMENTS.get(character, character)

        try:
            character.encode("cp1252")
        except UnicodeEncodeError:
            continue

        safe_characters.append(character)

    return "".join(safe_characters)


def _pdf_markup(runs: list[Run]) -> str:
    parts = []

    for run in runs:
        text = xml_escape(_pdf_safe(run.text)).replace("\n", "<br/>")

        if not text:
            continue

        if run.bold:
            text = f"<b>{text}</b>"

        if run.italic:
            text = f"<i>{text}</i>"

        if run.link:
            href = xml_escape(run.link, {'"': "&quot;"})
            text = f'<a href="{href}" color="#{OLIVE_DARK}">{text}</a>'

        parts.append(text)

    return "".join(parts)


def _pdf_styles() -> dict[str, ParagraphStyle]:
    text_color = colors.HexColor(f"#{TEXT}")
    heading_color = colors.HexColor(f"#{OLIVE_DARK}")

    body = ParagraphStyle(
        "body",
        fontName="Helvetica",
        fontSize=10.5,
        leading=15,
        textColor=text_color,
        spaceAfter=6,
        alignment=TA_LEFT,
    )

    return {
        "title": ParagraphStyle(
            "title",
            parent=body,
            fontName="Times-Bold",
            fontSize=22,
            leading=27,
            textColor=heading_color,
            spaceAfter=2,
        ),
        "subtitle": ParagraphStyle(
            "subtitle",
            parent=body,
            fontSize=9,
            textColor=colors.HexColor(f"#{MUTED}"),
            spaceAfter=10,
        ),
        "h1": ParagraphStyle(
            "h1",
            parent=body,
            fontName="Times-Bold",
            fontSize=17,
            leading=22,
            textColor=heading_color,
            spaceBefore=12,
            spaceAfter=6,
        ),
        "h2": ParagraphStyle(
            "h2",
            parent=body,
            fontName="Times-Bold",
            fontSize=14,
            leading=18,
            textColor=heading_color,
            spaceBefore=10,
            spaceAfter=4,
        ),
        "h3": ParagraphStyle(
            "h3",
            parent=body,
            fontName="Helvetica-Bold",
            fontSize=11.5,
            leading=15,
            textColor=colors.HexColor(f"#{OLIVE}"),
            spaceBefore=8,
            spaceAfter=3,
        ),
        "body": body,
        "cell": ParagraphStyle(
            "cell",
            parent=body,
            fontSize=9,
            leading=12,
            spaceAfter=0,
        ),
        "header_cell": ParagraphStyle(
            "header_cell",
            parent=body,
            fontName="Helvetica-Bold",
            fontSize=9,
            leading=12,
            spaceAfter=0,
            textColor=colors.white,
        ),
        "code": ParagraphStyle(
            "code",
            parent=body,
            fontName="Courier",
            fontSize=8.5,
            leading=11,
            backColor=colors.HexColor(f"#{IVORY}"),
        ),
    }


def _draw_pdf_page(canvas, document) -> None:
    width, height = A4
    canvas.saveState()

    logo_width = 42 * mm
    logo_height = logo_width * 193 / 642
    canvas.drawImage(
        ImageReader(str(LOGO_FILE)),
        document.leftMargin,
        height - 12 * mm - logo_height,
        width=logo_width,
        height=logo_height,
        mask="auto",
    )

    canvas.setStrokeColor(colors.HexColor(f"#{GOLD}"))
    canvas.setLineWidth(1.2)
    header_line_y = height - 16 * mm - logo_height
    canvas.line(
        document.leftMargin,
        header_line_y,
        width - document.rightMargin,
        header_line_y,
    )

    canvas.setStrokeColor(colors.HexColor(f"#{BORDER}"))
    canvas.setLineWidth(0.6)
    canvas.line(
        document.leftMargin,
        16 * mm,
        width - document.rightMargin,
        16 * mm,
    )

    # Rodapé: só o logo (que já traz o nome); número da página à direita.
    footer_logo_width = 20 * mm
    footer_logo_height = footer_logo_width * 193 / 642
    canvas.drawImage(
        ImageReader(str(LOGO_FILE)),
        document.leftMargin,
        8 * mm,
        width=footer_logo_width,
        height=footer_logo_height,
        mask="auto",
    )

    canvas.setFont("Helvetica", 7.5)
    canvas.setFillColor(colors.HexColor(f"#{MUTED}"))
    canvas.drawRightString(
        width - document.rightMargin,
        9.5 * mm,
        str(document.page),
    )

    canvas.restoreState()


def _pdf_table(block: Block, styles: dict, available_width: float) -> Table:
    column_count = max(len(row) for row in block.rows)
    data = []

    for row_index, row in enumerate(block.rows):
        style = styles["header_cell"] if row_index == 0 else styles["cell"]
        cells = [Paragraph(_pdf_markup(cell), style) for cell in row]
        cells += [""] * (column_count - len(cells))
        data.append(cells)

    table = Table(
        data,
        colWidths=[available_width / column_count] * column_count,
        repeatRows=1,
    )
    commands = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{OLIVE}")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor(f"#{BORDER}")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]

    for row_index in range(2, len(data), 2):
        commands.append(
            (
                "BACKGROUND",
                (0, row_index),
                (-1, row_index),
                colors.HexColor(f"#{IVORY}"),
            )
        )

    table.setStyle(TableStyle(commands))
    return table


def export_pdf(content: str, title: str) -> bytes:
    buffer = BytesIO()
    logo_height = 42 * mm * 193 / 642
    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
        topMargin=22 * mm + logo_height,
        bottomMargin=22 * mm,
        title=_pdf_safe(title),
        author=_pdf_safe(SIGNATURE),
        creator="ÁGORA - Latitudes",
    )
    styles = _pdf_styles()
    title, blocks = _split_title(parse_markdown(content), title)
    story = [
        Paragraph(xml_escape(_pdf_safe(title)), styles["title"]),
        Paragraph(f"Gerado em {_today()}", styles["subtitle"]),
    ]

    for block in blocks:
        if block.kind == "heading":
            style = styles[f"h{min(block.level, 3)}"]
            story.append(Paragraph(_pdf_markup(block.runs), style))
        elif block.kind == "paragraph":
            story.append(Paragraph(_pdf_markup(block.runs), styles["body"]))
        elif block.kind == "list_item":
            indent = 14 + block.depth * 14
            item_style = ParagraphStyle(
                f"item_{block.depth}",
                parent=styles["body"],
                leftIndent=indent,
                bulletIndent=indent - 11,
                spaceAfter=3,
            )
            story.append(
                Paragraph(
                    _pdf_markup(block.runs),
                    item_style,
                    bulletText=block.marker or None,
                )
            )
        elif block.kind == "table":
            story.append(Spacer(1, 4))
            story.append(_pdf_table(block, styles, document.width))
            story.append(Spacer(1, 8))
        elif block.kind == "rule":
            story.append(
                HRFlowable(
                    width="100%",
                    thickness=0.6,
                    color=colors.HexColor(f"#{BORDER}"),
                    spaceBefore=6,
                    spaceAfter=6,
                )
            )
        elif block.kind == "code":
            story.append(Preformatted(_pdf_safe(block.text), styles["code"]))

    document.build(
        story,
        onFirstPage=_draw_pdf_page,
        onLaterPages=_draw_pdf_page,
    )
    return buffer.getvalue()


# --------------------------------------------------------------------------
# Word (python-docx)
# --------------------------------------------------------------------------


def _docx_add_runs(paragraph, runs: list[Run]) -> None:
    for run in runs:
        text = run.text

        if run.link and run.link not in text:
            text = f"{text} ({run.link})"

        docx_run = paragraph.add_run(text)
        docx_run.bold = run.bold or None
        docx_run.italic = run.italic or None

        if run.link:
            docx_run.font.color.rgb = RGBColor.from_string(OLIVE_DARK)
            docx_run.underline = True


def _docx_shade(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), color)
    properties.append(shading)


def _docx_page_number(paragraph) -> None:
    run = paragraph.add_run()
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.append(field_begin)
    run._r.append(instruction)
    run._r.append(field_end)


def _docx_set_font(style, name: str, size: float, color: str) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(
        qn("w:eastAsia"),
        name,
    )


def export_docx(content: str, title: str) -> bytes:
    document = Document()
    section = document.sections[0]
    section.left_margin = section.right_margin = Cm(2)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2)

    _docx_set_font(document.styles["Normal"], "Arial", 10.5, TEXT)
    for level, size in ((1, 17), (2, 14), (3, 12)):
        heading_style = document.styles[f"Heading {level}"]
        _docx_set_font(
            heading_style,
            "Georgia",
            size,
            OLIVE_DARK if level < 3 else OLIVE,
        )
        heading_style.font.bold = True

    header_paragraph = section.header.paragraphs[0]
    header_paragraph.add_run().add_picture(str(LOGO_FILE), width=Cm(4.2))

    # Rodapé: só o logo (que já traz o nome); número da página à direita.
    footer = section.footer.paragraphs[0]
    footer.paragraph_format.tab_stops.add_tab_stop(
        section.page_width - section.left_margin - section.right_margin,
        alignment=2,
    )
    footer.add_run().add_picture(str(LOGO_FILE), width=Cm(2))
    footer.add_run("\t")
    _docx_page_number(footer)

    for run in footer.runs[-1:]:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MUTED)

    title, blocks = _split_title(parse_markdown(content), title)

    title_paragraph = document.add_paragraph()
    title_run = title_paragraph.add_run(title)
    title_run.bold = True
    title_run.font.name = "Georgia"
    title_run.font.size = Pt(22)
    title_run.font.color.rgb = RGBColor.from_string(OLIVE_DARK)

    date_paragraph = document.add_paragraph()
    date_run = date_paragraph.add_run(f"Gerado em {_today()}")
    date_run.font.size = Pt(9)
    date_run.font.color.rgb = RGBColor.from_string(MUTED)

    for block in blocks:
        if block.kind == "heading":
            paragraph = document.add_heading(level=min(block.level, 3))
            _docx_add_runs(paragraph, block.runs)
        elif block.kind == "paragraph":
            _docx_add_runs(document.add_paragraph(), block.runs)
        elif block.kind == "list_item":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.8 + block.depth * 0.7)
            paragraph.paragraph_format.first_line_indent = Cm(-0.5)
            paragraph.paragraph_format.space_after = Pt(2)
            prefix = f"{block.marker}\t" if block.marker else "\t"
            paragraph.add_run(prefix)
            _docx_add_runs(paragraph, block.runs)
        elif block.kind == "table":
            column_count = max(len(row) for row in block.rows)
            table = document.add_table(
                rows=len(block.rows),
                cols=column_count,
            )
            table.style = "Table Grid"

            for row_index, row in enumerate(block.rows):
                for column_index, cell_runs in enumerate(row):
                    cell = table.cell(row_index, column_index)
                    paragraph = cell.paragraphs[0]
                    _docx_add_runs(paragraph, cell_runs)

                    for run in paragraph.runs:
                        run.font.size = Pt(9)

                    if row_index == 0:
                        _docx_shade(cell, OLIVE)

                        for run in paragraph.runs:
                            run.bold = True
                            run.font.color.rgb = RGBColor.from_string("FFFFFF")
                    elif row_index % 2 == 0:
                        _docx_shade(cell, IVORY)

            document.add_paragraph()
        elif block.kind == "rule":
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rule_run = paragraph.add_run("—" * 20)
            rule_run.font.color.rgb = RGBColor.from_string(BORDER)
        elif block.kind == "code":
            paragraph = document.add_paragraph()
            code_run = paragraph.add_run(block.text)
            code_run.font.name = "Consolas"
            code_run.font.size = Pt(9)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# --------------------------------------------------------------------------
# Planilhas (Excel e CSV) — exportam as tabelas da resposta
# --------------------------------------------------------------------------

_NUMBER_PATTERN = re.compile(r"^-?(?:\d{1,3}(?:\.\d{3})+|\d+)(?:,\d+)?$")


def _cell_value(text: str):
    # Números no formato brasileiro viram números de verdade na planilha.
    if _NUMBER_PATTERN.match(text):
        number = float(text.replace(".", "").replace(",", "."))
        return int(number) if number.is_integer() else number

    return text


def _tables(content: str) -> list[list[list[str]]]:
    return [
        [[_plain_text(cell) for cell in row] for row in block.rows]
        for block in parse_markdown(content)
        if block.kind == "table"
    ]


def _xlsx_logo() -> XlsxImage:
    with PilImage.open(LOGO_FILE) as logo:
        resized = logo.copy()
        resized.thumbnail((210, 64))

    buffer = BytesIO()
    resized.save(buffer, format="PNG")
    buffer.seek(0)
    return XlsxImage(buffer)


def export_xlsx(content: str, title: str) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    thin = Side(style="thin", color=BORDER)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for table_index, rows in enumerate(_tables(content), start=1):
        sheet = workbook.create_sheet(f"Tabela {table_index}")
        sheet.sheet_view.showGridLines = False
        sheet.add_image(_xlsx_logo(), "A1")
        sheet.row_dimensions[1].height = 22
        sheet.row_dimensions[2].height = 22

        sheet["A4"] = title
        sheet["A4"].font = Font(name="Georgia", size=14, bold=True, color=OLIVE_DARK)
        sheet["A5"] = f"Gerado em {_today()}"
        sheet["A5"].font = Font(name="Arial", size=8, italic=True, color=MUTED)

        header_row = 7
        column_count = max(len(row) for row in rows)
        widths = [10] * column_count

        for row_offset, row in enumerate(rows):
            excel_row = header_row + row_offset

            for column_index in range(column_count):
                text = row[column_index] if column_index < len(row) else ""
                cell = sheet.cell(
                    row=excel_row,
                    column=column_index + 1,
                    value=text if row_offset == 0 else _cell_value(text),
                )
                cell.border = border
                cell.alignment = Alignment(wrap_text=True, vertical="top")

                if row_offset == 0:
                    cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor=OLIVE)
                else:
                    cell.font = Font(name="Arial", size=10, color=TEXT)

                    if row_offset % 2 == 0:
                        cell.fill = PatternFill("solid", fgColor=IVORY)

                widths[column_index] = max(
                    widths[column_index],
                    min(len(text) + 2, 60),
                )

        for column_index, width in enumerate(widths, start=1):
            sheet.column_dimensions[
                sheet.cell(row=1, column=column_index).column_letter
            ].width = width

        sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
        # O logo já está no topo da aba; o rodapé de impressão só numera.
        sheet.oddFooter.right.text = "&P"

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def export_csv(content: str) -> bytes:
    # Separador ";" e BOM UTF-8: o Excel em português abre direto, com acentos.
    output = StringIO()
    writer = csv.writer(output, delimiter=";")

    for table_index, rows in enumerate(_tables(content)):
        if table_index:
            writer.writerow([])

        writer.writerows(rows)

    return output.getvalue().encode("utf-8-sig")
